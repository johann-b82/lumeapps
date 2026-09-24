"""Produktions-Priorisierung — Rechenkern (v1.126).

Priorisiert wird auf **BA-Ebene** (Betriebsauftrag = Auftrag/Vorgang). Die
Positionen eines BA — die FAs der Fertigartikel samt ihrer Halbzeug-FAs —
laufen mit.

Reihenfolge der Gesamtliste:
  1. Je Position gilt der Termin der gewinnenden Prioliste, sonst das
     Lieferdatum aus dem Auftrag, und der Listenrang. Steht eine Position in
     mehreren Listen, gewinnt die Liste mit dem jüngsten Stand.
  2. Ein BA sortiert nach seinem frühesten Positionstermin, dann nach seinem
     besten (kleinsten) Listenrang, dann nach BA-Nummer.
  3. Die gespeicherte manuelle BA-Reihenfolge geht vor: manuell sortierte BAs
     behalten untereinander ihre Reihenfolge und belegen die Plätze, die sie
     automatisch hätten. Neue BAs landen so an ihrem automatischen Platz.
  4. Innerhalb eines BA: Termin, Rang, Pos, UPos.

Bereichslisten lösen jede Position über den Ressourcenplan in Tätigkeiten auf
(Komponenten rekursiv, in Positionsreihenfolge — so steht ein Halbzeug vor dem
Arbeitsgang, der es verbaut) und filtern sie nach Bereich. Zeiten sind in
Minuten: Rüstzeit einmal, Operativzeit je Stück.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    PrioArtikel,
    PrioAuftragPosition,
    PrioListe,
    PrioListeEintrag,
    PrioManuell,
    PrioPlanPosition,
)

MAX_TIEFE = 15
_DREI = Decimal("0.001")
_OHNE_RANG = 10**9

# Bereich → Ressourcen (Apollo RES-Nummern). Varianten wie 10000MA01 zählen
# zur Stammressource; alle 1xx000 sind Fremdvergabe.
BEREICHE: list[tuple[str, str, tuple[str, ...]]] = [
    ("zuschnitt", "Zuschnitt", ("10000", "16000")),
    ("naeherei", "Näherei", ("13000", "20000", "20002", "20003", "21000", "22000")),
    ("bezieherei", "Bezieherei", ("30000", "30002", "30003", "30004")),
    ("schaeumerei", "Schäumerei", ("31000", "31002", "31003")),
    ("wandverkleidung", "Wandverkleidung", ("40000", "41000", "48000")),
    ("teppiche", "Teppiche", ("40001", "40002")),
    ("spezialserie", "Spezialserie", ("50000", "50001")),
    ("gurte", "Gurte", ("60000",)),
    ("qs", "QS", ("70000", "71000")),
    ("verpackung", "Verpackung", ("80000", "81000")),
    ("fremdvergabe", "Fremdvergabe", ("1000", "1400")),
    ("bemusterung", "Bemusterung", ("90000", "90010")),
    ("etikettendruck", "Etikettendruck", ("11000",)),
    ("sonstige", "Sonstige", ()),
]
_RES_ZU_BEREICH = {r: key for key, _, res in BEREICHE for r in res}
BEREICH_LABEL = {key: label for key, label, _ in BEREICHE}


def bereich_fuer(ressource: str) -> str:
    r = re.sub(r"MA\d+$", "", ressource.strip())
    if r in _RES_ZU_BEREICH:
        return _RES_ZU_BEREICH[r]
    if re.fullmatch(r"1\d{2}000", r):
        return "fremdvergabe"
    return "sonstige"


Schluessel = tuple[str, int, int]


@dataclass
class ListenTreffer:
    liste_id: int
    liste_name: str
    stand: date
    rang: int | None
    termin: date | None
    kommentar: str | None


def waehle_treffer(treffer: list[ListenTreffer]) -> ListenTreffer | None:
    """Jüngster Stand gewinnt; bei gleichem Stand die später importierte Liste."""
    return max(treffer, key=lambda t: (t.stand, t.liste_id), default=None)


def _vorgang_sortwert(v: str) -> tuple[int, str]:
    return (int(v), "") if v.isdigit() else (10**12, v)


def _termin_rang(lieferdatum: date | None, t: ListenTreffer | None) -> tuple[date, int]:
    termin = (t.termin if t and t.termin else None) or lieferdatum
    return (termin or date.max, t.rang if t and t.rang is not None else _OHNE_RANG)


def ba_reihenfolge(
    positionen: dict[Schluessel, date | None],
    treffer: dict[Schluessel, ListenTreffer],
) -> list[str]:
    """Automatische BA-Reihenfolge: frühester Termin, bester Rang, BA-Nummer."""
    wert: dict[str, tuple[date, int]] = {}
    for k, lieferdatum in positionen.items():
        termin, rang = _termin_rang(lieferdatum, treffer.get(k))
        alt = wert.get(k[0])
        wert[k[0]] = (min(alt[0], termin), min(alt[1], rang)) if alt else (termin, rang)
    return sorted(wert, key=lambda v: (*wert[v], _vorgang_sortwert(v)))


def positions_reihenfolge(
    positionen: dict[Schluessel, date | None],
    treffer: dict[Schluessel, ListenTreffer],
) -> list[Schluessel]:
    """Reihenfolge innerhalb eines BA: Termin, Rang, Pos, UPos."""
    return sorted(positionen, key=lambda k: (*_termin_rang(positionen[k], treffer.get(k)), k[1], k[2]))


def wende_manuell_an(auto: list[str], manuell: dict[str, int]) -> list[str]:
    manuelle = sorted((k for k in auto if k in manuell), key=lambda k: manuell[k])
    it = iter(manuelle)
    return [next(it) if k in manuell else k for k in auto]


@dataclass
class PlanZeile:
    pos: int
    typ: str
    ref: str
    text: str | None
    menge: Decimal | None
    kostenstelle: str | None
    ruestzeit: Decimal | None
    operativzeit: Decimal | None


@dataclass
class Taetigkeit:
    ebene: int
    pfad: str
    artikelnr: str
    ressource: str
    bereich: str
    text: str | None
    kostenstelle: str | None
    menge: Decimal
    ruestzeit: Decimal
    operativzeit: Decimal


def loese_auf(artikelnr: str, plan: dict[str, list[PlanZeile]], menge: Decimal = Decimal(1)) -> list[Taetigkeit]:
    """Alle Arbeitsgänge eines Artikels über alle Ebenen, zyklensicher."""
    ergebnis: list[Taetigkeit] = []

    def gehe(art: str, faktor: Decimal, ebene: int, pfad: str, weg: frozenset[str]) -> None:
        for z in plan.get(art, ()):
            p = f"{pfad}.{z.pos}" if pfad else str(z.pos)
            if z.typ == "RES":
                ergebnis.append(Taetigkeit(
                    ebene=ebene, pfad=p, artikelnr=art, ressource=z.ref,
                    bereich=bereich_fuer(z.ref), text=z.text, kostenstelle=z.kostenstelle,
                    menge=faktor, ruestzeit=z.ruestzeit or Decimal(0),
                    operativzeit=z.operativzeit or Decimal(0),
                ))
            elif z.ref not in weg and ebene < MAX_TIEFE:
                gehe(z.ref, faktor * (z.menge if z.menge is not None else Decimal(1)),
                     ebene + 1, p, weg | {z.ref})

    gehe(artikelnr, menge, 0, "", frozenset({artikelnr}))
    return ergebnis


# ── Datenbank ──────────────────────────────────────────────────────────────


@dataclass
class PositionsZeile:
    pos: int
    upos: int
    artikelnr: str | None
    bezeichnung: str | None
    menge: Decimal | None
    einheit: str | None
    lieferdatum: date | None
    termin: date | None
    liste: str | None
    listen_rang: int | None
    kommentar: str | None
    gesperrt: bool
    ohne_plan: bool


@dataclass
class BaZeile:
    rang: int
    vorgang_nr: str
    kunde: str | None
    termin: date | None
    listen_rang: int | None
    gesperrt: bool
    manuell: bool
    positionen: list[PositionsZeile]


async def lade_gesamtliste(db: AsyncSession) -> list[BaZeile]:
    positionen = {
        (p.vorgang_nr, p.pos, p.upos): p
        for p in (await db.execute(select(PrioAuftragPosition))).scalars()
    }
    alle_treffer: dict[Schluessel, list[ListenTreffer]] = {}
    rows = await db.execute(
        select(PrioListeEintrag, PrioListe).join(PrioListe, PrioListe.id == PrioListeEintrag.liste_id)
    )
    for e, liste in rows.all():
        alle_treffer.setdefault((e.vorgang_nr, e.pos, e.upos), []).append(ListenTreffer(
            liste_id=liste.id, liste_name=liste.name, stand=liste.stand,
            rang=e.rang, termin=e.termin, kommentar=e.kommentar,
        ))
    treffer = {k: t for k, ts in alle_treffer.items() if (t := waehle_treffer(ts))}
    manuell = {m.vorgang_nr: m.reihenfolge for m in (await db.execute(select(PrioManuell))).scalars()}
    geplant = set((await db.execute(select(PrioPlanPosition.artikelnr).distinct())).scalars())

    lieferdaten = {k: p.lieferdatum for k, p in positionen.items()}
    je_ba: dict[str, dict[Schluessel, date | None]] = {}
    for k, d in lieferdaten.items():
        je_ba.setdefault(k[0], {})[k] = d

    zeilen: list[BaZeile] = []
    for i, ba in enumerate(wende_manuell_an(ba_reihenfolge(lieferdaten, treffer), manuell), start=1):
        pos_zeilen: list[PositionsZeile] = []
        for k in positions_reihenfolge(je_ba[ba], treffer):
            p, t = positionen[k], treffer.get(k)
            pos_zeilen.append(PositionsZeile(
                pos=k[1], upos=k[2], artikelnr=p.artikelnr, bezeichnung=p.bezeichnung,
                menge=p.menge, einheit=p.einheit, lieferdatum=p.lieferdatum,
                termin=(t.termin if t and t.termin else p.lieferdatum),
                liste=t.liste_name if t else None, listen_rang=t.rang if t else None,
                kommentar=t.kommentar if t else None, gesperrt=p.gesperrt,
                ohne_plan=bool(p.artikelnr) and p.artikelnr not in geplant,
            ))
        termine = [z.termin for z in pos_zeilen if z.termin]
        raenge = [z.listen_rang for z in pos_zeilen if z.listen_rang is not None]
        zeilen.append(BaZeile(
            rang=i, vorgang_nr=ba, kunde=positionen[next(iter(je_ba[ba]))].kunde,
            termin=min(termine, default=None), listen_rang=min(raenge, default=None),
            gesperrt=any(z.gesperrt for z in pos_zeilen), manuell=ba in manuell,
            positionen=pos_zeilen,
        ))
    return zeilen


async def lade_plan(db: AsyncSession) -> dict[str, list[PlanZeile]]:
    plan: dict[str, list[PlanZeile]] = {}
    rows = await db.execute(select(PrioPlanPosition).order_by(PrioPlanPosition.artikelnr, PrioPlanPosition.pos))
    for z in rows.scalars():
        plan.setdefault(z.artikelnr, []).append(PlanZeile(
            pos=z.pos, typ=z.typ, ref=z.ref, text=z.text, menge=z.menge,
            kostenstelle=z.kostenstelle, ruestzeit=z.ruestzeit, operativzeit=z.operativzeit,
        ))
    return plan


@dataclass
class BereichsZeile:
    prio: int
    vorgang_nr: str
    pos: int
    upos: int
    termin: date | None
    kunde: str | None
    endartikel: str | None
    endartikel_bez: str | None
    ebene: int
    pfad: str
    artikelnr: str
    artikel_bez: str | None
    ressource: str
    kostenstelle: str | None
    taetigkeit: str | None
    menge: Decimal
    minuten: Decimal
    gesperrt: bool
    kommentar: str | None


@dataclass
class Bereichsdaten:
    zeilen: dict[str, list[BereichsZeile]] = field(default_factory=dict)


async def lade_bereiche(db: AsyncSession, nur: str | None = None) -> Bereichsdaten:
    gesamt = await lade_gesamtliste(db)
    plan = await lade_plan(db)
    namen = dict((await db.execute(select(PrioArtikel.artikelnr, PrioArtikel.bezeichnung))).all())
    cache: dict[str, list[Taetigkeit]] = {}
    daten = Bereichsdaten(zeilen={key: [] for key, _, _ in BEREICHE})
    for ba in gesamt:
        for p in ba.positionen:
            if not p.artikelnr:
                continue
            if p.artikelnr not in cache:
                cache[p.artikelnr] = loese_auf(p.artikelnr, plan)
            menge = p.menge if p.menge is not None else Decimal(1)
            for t in cache[p.artikelnr]:
                if nur and t.bereich != nur:
                    continue
                m = (t.menge * menge).quantize(_DREI)
                daten.zeilen[t.bereich].append(BereichsZeile(
                    prio=ba.rang, vorgang_nr=ba.vorgang_nr, pos=p.pos, upos=p.upos, termin=p.termin,
                    kunde=ba.kunde, endartikel=p.artikelnr, endartikel_bez=p.bezeichnung,
                    ebene=t.ebene, pfad=t.pfad, artikelnr=t.artikelnr, artikel_bez=namen.get(t.artikelnr),
                    ressource=t.ressource, kostenstelle=t.kostenstelle, taetigkeit=t.text,
                    menge=m, minuten=(t.ruestzeit + t.operativzeit * m).quantize(_DREI),
                    gesperrt=p.gesperrt, kommentar=p.kommentar,
                ))
    return daten


def bereich_xlsx(label: str, zeilen: list[BereichsZeile]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = label[:31]
    kopf = ["Prio", "BA", "Pos", "UPos", "Termin", "Kunde", "Endartikel", "Bezeichnung",
            "Ebene", "Artikel", "Artikelbezeichnung", "Ressource", "Kostenstelle",
            "Tätigkeit", "Menge", "Minuten", "Gesperrt", "Kommentar"]
    ws.append(kopf)
    for c in ws[1]:
        c.font = Font(bold=True)
    for z in zeilen:
        ws.append([z.prio, z.vorgang_nr, z.pos, z.upos, z.termin, z.kunde, z.endartikel,
                   z.endartikel_bez, z.ebene, z.artikelnr, z.artikel_bez, z.ressource,
                   z.kostenstelle, (z.taetigkeit or "").replace("\n", " "), float(z.menge),
                   round(float(z.minuten), 2), "J" if z.gesperrt else "", z.kommentar])
    for zelle in ws["E"][1:]:
        zelle.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
