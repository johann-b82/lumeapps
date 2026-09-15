"""Parser für die Produktions-Priorisierung (v1.126).

Drei Apollo-Exporte und eine Kunden-Prioritätsliste:

  * ``AswStm.txt``      Artikelstamm — Tab-getrennt, Latin-1. Enthält vereinzelt
                        Excel-Formelreste (``="..."``), die die Spaltenzahl
                        sprengen; solche Zeilen werden übersprungen.
  * ``dev_excel.txt``   RESSOURCENPLAN — 3 Kopfzeilen, dann eine Zeile je
                        Position eines Artikels (Komponente, Material oder
                        Arbeitsgang). Übernommen werden nur HLB/ART-Komponenten
                        und RES-Arbeitsgänge.
  * ``AswKpf.txt``      Auftragspositionen — über den bestehenden
                        ``auftrag_positionen_parser``.
  * Diehl-LoB (.xlsx)   Blatt ``Prioliste_Diehl`` (Rang je MSN) + Blatt
                        ``Auftragsbestand_ACM`` (Code1 → Auftrag/Pos/UP).

Die Funktionen lesen bewusst zeilenweise statt über pandas: die Exporte
enthalten unbalancierte Anführungszeichen, an denen der CSV-Parser scheitert.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.parsing.auftrag_positionen_parser import parse_auftrag_positionen_file
from app.security.limits import pruefe_archivgroesse

MAX_WARNUNGEN = 50


@dataclass
class Ergebnis:
    zeilen: list[dict[str, Any]] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)
    warnungen_anzahl: int = 0

    def warne(self, text: str) -> None:
        self.warnungen_anzahl += 1
        if len(self.warnungen) < MAX_WARNUNGEN:
            self.warnungen.append(text)


def _dekodiere(daten: bytes) -> list[str]:
    try:
        text = daten.decode("utf-8")
    except UnicodeDecodeError:
        text = daten.decode("latin-1")
    return text.splitlines()


def _s(val: Any) -> str:
    return "" if val is None else str(val).strip()


def _dezimal(val: Any) -> Decimal | None:
    s = _s(val)
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


# ── Artikelstamm ───────────────────────────────────────────────────────────

_STM_KOPF = ["Typ", "Artikelnr", "Bezeichnung 1", "Bezeichnung 2", "Einh"]


def parse_artikelstamm(daten: bytes) -> Ergebnis:
    zeilen = _dekodiere(daten)
    if not zeilen or zeilen[0].split("\t")[: len(_STM_KOPF)] != _STM_KOPF:
        raise ValueError("Kein Artikelstamm (erwartet Kopfzeile Typ, Artikelnr, Bezeichnung 1 …).")
    breite = zeilen[0].count("\t")
    erg = Ergebnis()
    gesehen: set[str] = set()
    for nr, zeile in enumerate(zeilen[1:], start=2):
        if not zeile.strip("\t "):
            continue
        felder = zeile.split("\t")
        if zeile.count("\t") != breite or '"' in felder[1]:
            erg.warne(f"Zeile {nr}: beschädigt (Spaltenzahl), übersprungen")
            continue
        artikelnr = _s(felder[1])
        if not artikelnr or artikelnr in gesehen:
            if artikelnr:
                erg.warne(f"Zeile {nr}: Artikel {artikelnr} doppelt, erste Zeile gilt")
            continue
        gesehen.add(artikelnr)
        erg.zeilen.append({
            "artikelnr": artikelnr,
            "typ": _s(felder[0])[:8],
            "bezeichnung": _s(felder[2])[:255] or None,
            "bezeichnung2": _s(felder[3])[:255] or None,
            "einheit": _s(felder[4])[:16] or None,
        })
    return erg


# ── Ressourcenplan ─────────────────────────────────────────────────────────

# Spaltenpositionen im RESSOURCENPLAN (Kopfzeile in Zeile 4, Namen doppelt).
_RP_KOPF = {0: "Artikel", 9: "RSC-Pos", 10: "Kostenstelle", 12: "Ressource/Material",
            13: "Typ", 14: "Bez.1", 21: "Menge", 22: "Einheit", 23: "Rüstzeit",
            24: "Operativzeit"}
_RP_UEBERNEHMEN = {"HLB", "ART", "RES"}


def parse_ressourcenplan(daten: bytes) -> Ergebnis:
    zeilen = _dekodiere(daten)
    kopf_idx = next((i for i, z in enumerate(zeilen[:10]) if z.startswith("Artikel\t")), None)
    if kopf_idx is None or not zeilen[0].startswith("RESSOURCENPLAN"):
        raise ValueError("Kein Ressourcenplan (erwartet Titel RESSOURCENPLAN und Kopfzeile Artikel …).")
    kopf = zeilen[kopf_idx].split("\t")
    for i, name in _RP_KOPF.items():
        if i >= len(kopf) or kopf[i].strip() != name:
            raise ValueError(f"Ressourcenplan: Spalte {i + 1} heißt nicht „{name}“.")
    breite = zeilen[kopf_idx].count("\t")
    erg = Ergebnis()
    for nr, zeile in enumerate(zeilen[kopf_idx + 1:], start=kopf_idx + 2):
        if not zeile.strip("\t "):
            continue
        if zeile.count("\t") != breite:
            erg.warne(f"Zeile {nr}: beschädigt (Spaltenzahl), übersprungen")
            continue
        f = zeile.split("\t")
        typ = _s(f[13])
        if typ not in _RP_UEBERNEHMEN:
            continue
        artikelnr, ref = _s(f[0]), _s(f[12])
        try:
            pos = int(_s(f[9]))
        except ValueError:
            erg.warne(f"Zeile {nr}: RSC-Pos „{_s(f[9])}“ unlesbar, übersprungen")
            continue
        if not artikelnr or not ref:
            continue
        text = "\n".join(t for t in (_s(x) for x in f[14:20]) if t)
        erg.zeilen.append({
            "artikelnr": artikelnr[:64],
            "pos": pos,
            "typ": typ,
            "ref": ref[:64],
            "text": text or None,
            "menge": _dezimal(f[21]),
            "einheit": _s(f[22])[:16] or None,
            "kostenstelle": _s(f[10])[:16] or None,
            "ruestzeit": _dezimal(f[23]) if typ == "RES" else None,
            "operativzeit": _dezimal(f[24]) if typ == "RES" else None,
        })
    return erg


# ── Auftragspositionen ─────────────────────────────────────────────────────


def parse_auftraege(daten: bytes, dateiname: str) -> Ergebnis:
    rows, errors = parse_auftrag_positionen_file(daten, dateiname)
    if not rows and errors and errors[0].get("row") == 0:
        raise ValueError(f"Keine Auftragspositionen: {errors[0]['message']}")
    erg = Ergebnis()
    for e in errors:
        erg.warne(f"Zeile {e['row']}: {e['message']}")
    for r in rows:
        raw = r["raw"]
        erg.zeilen.append({
            "vorgang_nr": r["vorgang_nr"],
            "pos": r["pos"],
            "upos": r["upos"],
            "artikelnr": r["article_number"],
            "bezeichnung": (r["article_name"] or "")[:255] or None,
            "menge": r["quantity"],
            "einheit": r["unit"],
            "lieferdatum": r["lieferdatum"],
            "kunde_nr": r["customer_id"],
            "kunde": r["customer_name"],
            "fremdnr": r["external_order_nr"],
            "pos_typ_2": r["pos_typ_2"],
            "status": (raw.get("St") or None),
            "gesperrt": raw.get("Sperre manuell", "").upper() == "J",
        })
    return erg


# ── Diehl-Prioliste ────────────────────────────────────────────────────────


@dataclass
class ListenErgebnis(Ergebnis):
    stand: date | None = None


def _kopf_index(kopfzeile: tuple, pflicht: list[str], blatt: str) -> dict[str, int]:
    # Erste Spalte gewinnt: Auftragsbestand_ACM hat „Pos“ zweimal (K = Auftragsposition, AS = Vater-Pos).
    namen: dict[str, int] = {}
    for i, v in enumerate(kopfzeile):
        namen.setdefault(" ".join(_s(v).split()), i)
    fehlt = [p for p in pflicht if p not in namen]
    if fehlt:
        raise ValueError(f"Blatt „{blatt}“: Spalte(n) fehlen: {', '.join(fehlt)}")
    return namen


def _ganzzahl(val: Any) -> int | None:
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return int(val)
    s = _s(val)
    return int(s) if s.isdigit() else None


def _datum(val: Any) -> date | None:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def stand_aus_dateiname(dateiname: str) -> date | None:
    """``26-09-14_LoB_….xlsx`` → 2026-09-14."""
    m = re.match(r"^(\d{2})-(\d{2})-(\d{2})", dateiname or "")
    if not m:
        return None
    try:
        return date(2000 + int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def parse_diehl_prioliste(daten: bytes, dateiname: str) -> ListenErgebnis:
    pruefe_archivgroesse(daten)
    wb = load_workbook(io.BytesIO(daten), read_only=True, data_only=True)
    try:
        for blatt in ("Prioliste_Diehl", "Auftragsbestand_ACM"):
            if blatt not in wb.sheetnames:
                raise ValueError(f"Keine Diehl-Prioliste: Blatt „{blatt}“ fehlt.")

        # Code1 → ACM-Auftragspositionen
        ab = wb["Auftragsbestand_ACM"].iter_rows(values_only=True)
        k = _kopf_index(next(ab), ["Code1", "Nummer", "Pos", "UP"], "Auftragsbestand_ACM")
        code_zu_pos: dict[str, set[tuple[str, int, int]]] = {}
        for z in ab:
            code, nummer, pos = _s(z[k["Code1"]]), _ganzzahl(z[k["Nummer"]]), _ganzzahl(z[k["Pos"]])
            if code and nummer is not None and pos is not None:
                code_zu_pos.setdefault(code, set()).add((str(nummer), pos, _ganzzahl(z[k["UP"]]) or 0))

        pl = wb["Prioliste_Diehl"].iter_rows(values_only=True)
        k = _kopf_index(next(pl), ["Code01", "Reihenfolge", "ACM Liefertermin", "Kommentar ACM"],
                        "Prioliste_Diehl")
        erg = ListenErgebnis(stand=stand_aus_dateiname(dateiname))
        je_pos: dict[tuple[str, int, int], dict[str, Any]] = {}
        ohne_auftrag = 0
        for z in pl:
            code = _s(z[k["Code01"]])
            if not code:
                continue
            positionen = code_zu_pos.get(code)
            if not positionen:
                ohne_auftrag += 1
                continue
            for vorgang_nr, pos, upos in sorted(positionen):
                neu = {
                    "vorgang_nr": vorgang_nr, "pos": pos, "upos": upos,
                    "rang": _ganzzahl(z[k["Reihenfolge"]]),
                    "termin": _datum(z[k["ACM Liefertermin"]]),
                    "kommentar": _s(z[k["Kommentar ACM"]]) or None,
                    "extern_ref": code[:64],
                }
                alt = je_pos.get((vorgang_nr, pos, upos))
                # Mehrere Diehl-Zeilen auf derselben Position: kleinster Rang gilt.
                if alt is None or (neu["rang"] is not None and (alt["rang"] is None or neu["rang"] < alt["rang"])):
                    je_pos[(vorgang_nr, pos, upos)] = neu
        if ohne_auftrag:
            erg.warne(f"{ohne_auftrag} Zeilen ohne ACM-Auftragsposition (k. Auftrag / nicht erfasst) übersprungen")
        erg.zeilen = list(je_pos.values())
        return erg
    finally:
        wb.close()
