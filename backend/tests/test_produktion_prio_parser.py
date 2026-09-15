"""Parser-Tests für app.parsing.produktion_prio_parser (v1.126).

Die Fixtures bilden die echten Apollo-Exporte nach: Tab-getrennt, Latin-1,
deutsche Dezimalzahlen, samt der Excel-Formelreste, die die Spaltenzahl sprengen.
"""
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.parsing.produktion_prio_parser import (
    parse_artikelstamm,
    parse_auftraege,
    parse_diehl_prioliste,
    parse_ressourcenplan,
    stand_aus_dateiname,
)

STM_KOPF = "Typ\tArtikelnr\tBezeichnung 1\tBezeichnung 2\tEinh\tWGR\tGewicht\tNettogewicht\t\t\r\n"


def test_artikelstamm_liest_und_ueberspringt_defekte_zeilen():
    daten = (
        STM_KOPF
        + "RES\t41000\tWandverkleidungen A350\t\t\t\t\t\t\t\r\n"
        + "HLB\tH3058\tSec. Lining\tIndex A\tSTK\tA350\t0\t0\t\t\r\n"
        + 'ART\t="001052"\t"=""SEAT TPL"\t\t\tSTK\t\t0\t0\t\t\t\r\n'
        + "HLB\tH3058\tdoppelt\t\tSTK\t\t0\t0\t\t\r\n"
    ).encode("latin-1")
    erg = parse_artikelstamm(daten)
    assert [z["artikelnr"] for z in erg.zeilen] == ["41000", "H3058"]
    assert erg.zeilen[1] == {"artikelnr": "H3058", "typ": "HLB", "bezeichnung": "Sec. Lining",
                             "bezeichnung2": "Index A", "einheit": "STK"}
    assert erg.warnungen_anzahl == 2


def test_artikelstamm_falsche_datei():
    with pytest.raises(ValueError):
        parse_artikelstamm(b"Datum\tZeit\r\n")


def _rp_zeile(artikel, pos, kst, ref, typ, bez1, menge="", einheit="", ruest="", op=""):
    felder = [artikel, "ART", "Bez", "", "", "", "", "", "WG", pos, kst, "", ref, typ,
              bez1, "", "", "", "", "", "", menge, einheit, ruest, op, ""]
    return "\t".join(felder) + "\r\n"


RP_KOPF = (
    "RESSOURCENPLAN\tMandant 1" + "\t" * 24 + "\r\n"
    "15.09.2026\t08:20:23" + "\t" * 24 + "\r\n"
    + "\t" * 25 + "\r\n"
    "Artikel\tTyp\tBez.1\tBez.2\tBez.3\tBez.4\tBez.5\tBez.6\tWGrp\tRSC-Pos\tKostenstelle\t\t"
    "Ressource/Material\tTyp\tBez.1\tBez.2\tBez.3\tBez.4\tBez.5\tBez.6\tLohngruppe\tMenge\t"
    "Einheit\tRüstzeit\tOperativzeit\t\r\n"
)


def test_ressourcenplan_nimmt_komponenten_und_arbeitsgaenge():
    daten = (
        RP_KOPF
        + _rp_zeile("6817", "1", "", "H3058", "HLB", "Sec. Lining", "1", "STK", "0,00000000")
        + _rp_zeile("6817", "2", "", "L 3285", "MAT", "Hakenband", "2,9", "LFM", "0,00000000")
        + _rp_zeile("6817", "3", "4100", "41000", "RES", "- HB anbringen", ruest="0", op="20")
        + _rp_zeile("6817", "4", "7000", "70000", "RES", "- QW", ruest="1.000,5", op="15")
    ).encode("latin-1")
    erg = parse_ressourcenplan(daten)
    assert [(z["pos"], z["typ"], z["ref"]) for z in erg.zeilen] == [
        (1, "HLB", "H3058"), (3, "RES", "41000"), (4, "RES", "70000")]
    assert erg.zeilen[0]["menge"] == Decimal("1")
    assert erg.zeilen[0]["ruestzeit"] is None
    assert erg.zeilen[1]["kostenstelle"] == "4100"
    assert erg.zeilen[1]["operativzeit"] == Decimal("20")
    assert erg.zeilen[2]["ruestzeit"] == Decimal("1000.5")


def test_ressourcenplan_falsche_datei():
    with pytest.raises(ValueError):
        parse_ressourcenplan(b"ARBEITSFOLGENPLAN\tMandant 1\r\n")


def test_auftraege_uebernimmt_status_und_sperre():
    kopf = ("Typ\tVorgang Nr.\tPos\tUPos\tDatum\tAdr Nr.\tName 1\tOrt\tArtnr\tVersion\t"
            "Bezeichnung 1\tMenge\tME\tSt\tLieferdatum\tPreis\tPos Wert\tPos Typ 2\t"
            "Fremdnr\tSperre manuell\tSperre K-Limit\r\n")
    zeile = ("AUF\t1024906\t42\t0\t16.03.2026\t10005\tDiehl\tLaupheim\t12979\t\tHEAD PADDING\t"
             "1\tSTK\t1\t06.08.2026\t866,47\t819,74\tAV-F\tVR11S1024000000\tJ\t0\r\n")
    erg = parse_auftraege((kopf + zeile).encode("latin-1"), "AswKpf.txt")
    z = erg.zeilen[0]
    assert (z["vorgang_nr"], z["pos"], z["upos"], z["artikelnr"]) == ("1024906", 42, 0, "12979")
    assert z["lieferdatum"] == date(2026, 8, 6)
    assert z["status"] == "1" and z["gesperrt"] is True


def _diehl_xlsx() -> bytes:
    wb = Workbook()
    ab = wb.active
    ab.title = "Auftragsbestand_ACM"
    # Wie im Original: eine zweite Spalte „Pos“ (Vater-Pos) weiter rechts.
    ab.append(["Code1", "Wunsch-liefer\nKW", "Vor", "Nummer", "Pos", "UP", "Artikel", "Vater-Nr", "Pos"])
    ab.append(["4501139139500", 22.26, "AUF", 1024906, 50, 0, 7934, 0, -12])
    ab.append(["4501139139420", 22.26, "AUF", 1024906, 42, 0, 12979, 0, -11])
    pl = wb.create_sheet("Prioliste_Diehl")
    pl.append(["Code01", "ACM-BA-Nr.", "Reihenfolge", "ACM Liefertermin", "Kommentar ACM"])
    pl.append(["4501139139500", 1024906, 30, datetime(2026, 8, 6), None])
    pl.append(["4501139139420", 1024906, None, "geliefert", "Klärfall, Reklamation"])
    pl.append(["450115419410", "geliefert", 11, "geliefert", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_diehl_prioliste_ordnet_ueber_auftragsbestand_zu():
    erg = parse_diehl_prioliste(_diehl_xlsx(), "26-09-14_LoB_Auftragsbestand_Diehl_V02.xlsx")
    assert erg.stand == date(2026, 9, 14)
    by_pos = {z["pos"]: z for z in erg.zeilen}
    assert by_pos[50]["rang"] == 30 and by_pos[50]["termin"] == date(2026, 8, 6)
    assert by_pos[42]["rang"] is None and by_pos[42]["termin"] is None
    assert by_pos[42]["kommentar"] == "Klärfall, Reklamation"
    assert erg.warnungen_anzahl == 1  # Zeile ohne ACM-Auftrag


def test_diehl_prioliste_ohne_blatt():
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        parse_diehl_prioliste(buf.getvalue(), "x.xlsx")


def test_stand_aus_dateiname():
    assert stand_aus_dateiname("26-09-14_LoB.xlsx") == date(2026, 9, 14)
    assert stand_aus_dateiname("LoB.xlsx") is None
